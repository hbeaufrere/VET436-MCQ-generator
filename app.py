import os
import json
import re
import sqlite3
import hashlib
from pathlib import Path

from flask import Flask, render_template, request, jsonify, g
from dotenv import load_dotenv
import anthropic
import pdfplumber
from pptx import Presentation

load_dotenv()

app = Flask(__name__)
app.secret_key = os.getenv("SECRET_KEY", "dev-secret-change-me")

COURSE_MATERIALS = Path(__file__).parent / "course_materials"
COURSE_MATERIALS.mkdir(exist_ok=True)

ALLOWED_EXTENSIONS = {"pdf", "pptx", "ppt", "txt"}

DATABASE = Path(__file__).parent / "mcq.db"

# ---------------------------------------------------------------------------
# Database helpers
# ---------------------------------------------------------------------------

def get_db():
    if "db" not in g:
        g.db = sqlite3.connect(str(DATABASE))
        g.db.row_factory = sqlite3.Row
    return g.db


@app.teardown_appcontext
def close_db(exc):
    db = g.pop("db", None)
    if db is not None:
        db.close()


def init_db():
    db = sqlite3.connect(str(DATABASE))
    db.executescript("""
        CREATE TABLE IF NOT EXISTS documents (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            filename TEXT NOT NULL,
            file_hash TEXT NOT NULL UNIQUE,
            text_content TEXT NOT NULL,
            uploaded_at TIMESTAMP DEFAULT CURRENT_TIMESTAMP
        );
        CREATE TABLE IF NOT EXISTS quiz_sessions (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            questions_json TEXT NOT NULL,
            score INTEGER,
            total INTEGER,
            created_at TIMESTAMP DEFAULT CURRENT_TIMESTAMP
        );
    """)
    db.close()


init_db()

# ---------------------------------------------------------------------------
# Document parsing
# ---------------------------------------------------------------------------

def extract_text_pdf(filepath):
    text_parts = []
    with pdfplumber.open(filepath) as pdf:
        for page in pdf.pages:
            page_text = page.extract_text()
            if page_text:
                text_parts.append(page_text)
    return "\n\n".join(text_parts)


def extract_text_pptx(filepath):
    prs = Presentation(filepath)
    text_parts = []
    for slide in prs.slides:
        slide_texts = []
        for shape in slide.shapes:
            if shape.has_text_frame:
                for paragraph in shape.text_frame.paragraphs:
                    line = paragraph.text.strip()
                    if line:
                        slide_texts.append(line)
        if slide_texts:
            text_parts.append("\n".join(slide_texts))
    return "\n\n".join(text_parts)


def extract_text_txt(filepath):
    with open(filepath, "r", encoding="utf-8", errors="replace") as f:
        return f.read()


def extract_text(filepath):
    ext = Path(filepath).suffix.lower()
    if ext == ".pdf":
        return extract_text_pdf(filepath)
    elif ext in (".pptx", ".ppt"):
        return extract_text_pptx(filepath)
    elif ext == ".txt":
        return extract_text_txt(filepath)
    else:
        raise ValueError(f"Unsupported file type: {ext}")


def file_hash(filepath):
    h = hashlib.sha256()
    with open(filepath, "rb") as f:
        for chunk in iter(lambda: f.read(8192), b""):
            h.update(chunk)
    return h.hexdigest()

# ---------------------------------------------------------------------------
# Auto-seed: import all files from course_materials/ on startup
# ---------------------------------------------------------------------------

def seed_course_materials():
    """Scan course_materials/ and import any new files into the database."""
    db = sqlite3.connect(str(DATABASE))
    db.row_factory = sqlite3.Row
    count = 0

    for filepath in sorted(COURSE_MATERIALS.iterdir()):
        if filepath.is_file() and filepath.suffix.lower().lstrip(".") in ALLOWED_EXTENSIONS:
            try:
                fhash = file_hash(str(filepath))
                # Skip if already imported
                existing = db.execute(
                    "SELECT id FROM documents WHERE file_hash = ?", (fhash,)
                ).fetchone()
                if existing:
                    continue

                text = extract_text(str(filepath))
                if not text.strip():
                    print(f"  Skipped {filepath.name} (no text extracted)")
                    continue

                db.execute(
                    "INSERT INTO documents (filename, file_hash, text_content) VALUES (?, ?, ?)",
                    (filepath.name, fhash, text),
                )
                db.commit()
                count += 1
                print(f"  Imported: {filepath.name}")
            except Exception as e:
                print(f"  Error importing {filepath.name}: {e}")

    db.close()
    if count:
        print(f"Seeded {count} new document(s) from course_materials/")
    else:
        print("No new documents to import from course_materials/")


seed_course_materials()

# ---------------------------------------------------------------------------
# Claude API – MCQ generation
# ---------------------------------------------------------------------------

def sanitize_text(text):
    """Remove control characters and null bytes that break the API."""
    # Remove null bytes and other control chars (keep newlines, tabs, carriage returns)
    text = re.sub(r'[\x00-\x08\x0b\x0c\x0e-\x1f\x7f]', '', text)
    # Replace surrogate pairs / invalid unicode
    text = text.encode('utf-8', errors='replace').decode('utf-8', errors='replace')
    return text


def generate_mcqs(text_content, num_questions=10, topic_focus=None):
    """Call Claude to generate MCQs from document text."""
    client = anthropic.Anthropic()

    focus_instruction = ""
    if topic_focus:
        focus_instruction = f"\nFocus the questions specifically on: {topic_focus}"

    # Sanitize text to remove control characters from PDF extraction
    text_content = sanitize_text(text_content)

    # Truncate very long documents to fit context
    max_chars = 80_000
    if len(text_content) > max_chars:
        text_content = text_content[:max_chars] + "\n\n[Content truncated...]"

    prompt = f"""You are an expert veterinary science educator specializing in small mammal medicine (rabbits, guinea pigs, ferrets, chinchillas, hedgehogs, and other exotic small mammals). You are creating exam-style multiple choice questions (MCQs) for veterinary students.

Based on the following lecture/course material on small mammal medicine, generate exactly {num_questions} high-quality MCQs. All questions must be relevant to small mammal medicine.{focus_instruction}

Question type distribution (approximate):
- 25% Recall/knowledge questions: straightforward factual recall (e.g., anatomy, normal values, definitions)
- 50% Clinical scenario questions: present a patient case or clinical situation and ask for the best diagnosis, treatment, or next step
- 25% Comparative/species-differentiation questions: highlight differences between species (e.g., "Which species is the exception...", "How does X differ between ferrets and rabbits?")

Requirements for each question:
- Write a clear, specific question stem
- Provide exactly 4 answer options labeled A, B, C, D
- Exactly one option must be correct
- Include plausible distractors that test understanding, not just recall
- After the correct answer, provide a brief explanation (2-3 sentences) of WHY the correct answer is right and why key distractors are wrong
- Do NOT include questions about specific drug dosages, drug doses, or numerical blood/lab values (e.g., no "What is the normal blood glucose range..." or "What dose of meloxicam...")

Return your response as a JSON array with this exact structure:
[
  {{
    "question": "The question text",
    "options": {{
      "A": "First option",
      "B": "Second option",
      "C": "Third option",
      "D": "Fourth option"
    }},
    "correct_answer": "A",
    "explanation": "Explanation of why A is correct and why other options are incorrect."
  }}
]

Return ONLY the JSON array, no other text.

--- COURSE MATERIAL ---
{text_content}
"""

    message = client.messages.create(
        model="claude-haiku-4-5-20251001",
        max_tokens=4096,
        messages=[{"role": "user", "content": prompt}],
    )

    response_text = message.content[0].text.strip()

    # Calculate cost (Haiku 4.5: $1/M input, $5/M output)
    input_tokens = message.usage.input_tokens
    output_tokens = message.usage.output_tokens
    cost = (input_tokens / 1_000_000) * 1.0 + (output_tokens / 1_000_000) * 5.0

    # Extract JSON from response – handle markdown code blocks and preamble
    # Find the first '[' and last ']' to extract the JSON array
    start = response_text.find("[")
    end = response_text.rfind("]")
    if start != -1 and end != -1 and end > start:
        response_text = response_text[start:end + 1]

    return json.loads(response_text), cost

# ---------------------------------------------------------------------------
# Routes
# ---------------------------------------------------------------------------

@app.route("/")
def index():
    db = get_db()
    doc_count = db.execute("SELECT COUNT(*) as cnt FROM documents").fetchone()["cnt"]
    return render_template("index.html", doc_count=doc_count)


@app.route("/api/generate", methods=["POST"])
def api_generate():
    """Generate MCQs from all loaded course materials."""
    data = request.get_json()
    num_questions = min(int(data.get("num_questions", 10)), 15)
    topic_focus = (data.get("topic_focus") or "").strip() or None

    db = get_db()
    rows = db.execute("SELECT text_content, filename FROM documents").fetchall()

    if not rows:
        return jsonify({"error": "No course materials loaded. Place files in course_materials/ and restart."}), 400

    # Combine text from all documents
    combined_text = "\n\n".join(
        f"--- {row['filename']} ---\n{row['text_content']}" for row in rows
    )

    try:
        questions, cost = generate_mcqs(combined_text, num_questions, topic_focus)
    except json.JSONDecodeError as e:
        print(f"JSON parse error: {e}")
        return jsonify({"error": "Failed to parse generated questions. Please try again."}), 500
    except anthropic.APIError as e:
        return jsonify({"error": f"AI service error: {e.message}"}), 502
    except Exception as e:
        print(f"Unexpected error: {type(e).__name__}: {e}")
        return jsonify({"error": str(e)}), 500

    # Save quiz session
    db.execute(
        "INSERT INTO quiz_sessions (questions_json, total) VALUES (?, ?)",
        (json.dumps(questions), len(questions)),
    )
    db.commit()

    return jsonify({"questions": questions, "cost": round(cost, 4)})


# ---------------------------------------------------------------------------
# Canvas LTI / Embedding support
# ---------------------------------------------------------------------------

@app.after_request
def allow_iframe_embedding(response):
    """Allow the app to be embedded in Canvas iframes."""
    response.headers.pop("X-Frame-Options", None)
    response.headers["Content-Security-Policy"] = "frame-ancestors *"
    return response


if __name__ == "__main__":
    app.run(debug=True, host="0.0.0.0", port=5000)
